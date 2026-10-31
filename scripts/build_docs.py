#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_docs.py — 把课程仓库全部教学 md 转换为 Word（.docx）与 PDF 各一套。

规则（依据用户需求：除代码、PPT 外的 md 文件转换为 Word 和 PDF 各一套）：
  1. 收集 docs/ 下全部 .md + 根 README.md（跳过 07-项目资源 下同步自 dap-v2 的
     工程文档快照与两份 PDF 参考资料——它们是证据快照，教学分发以 Word/PDF 为主；
     如需包含，传 --with-project-docs）。
  2. 每份 md 生成同名 .docx 到 export/word/；文件名冲突时加父目录前缀。
  3. 用 LibreOffice headless 把全部 docx 转 PDF 到 export/pdf/。

用法：
  python3 scripts/build_docs.py [--with-project-docs] [--skip-pdf]
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
WORD_OUT = ROOT / "export" / "word"
PDF_OUT = ROOT / "export" / "pdf"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

C_PRIMARY = RGBColor(0x1B, 0x2A, 0x4A)
C_ACCENT = RGBColor(0x2E, 0x86, 0xDE)
C_MUTED = RGBColor(0x6B, 0x7A, 0x90)
C_CODE_BG = "F0F3F8"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def set_ea_font(target, ea_name):
    """给样式/run 设置中文字体（target 可为 style 对象或 run）。"""
    el = target.element if hasattr(target, "element") else target._r
    rPr = el.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), ea_name)


def style_setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0x2C, 0x33, 0x45)
    set_ea_font(st, "宋体")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (17, C_PRIMARY, "黑体", 12, 6),
        "Heading 2": (14, C_PRIMARY, "黑体", 10, 5),
        "Heading 3": (12, C_ACCENT, "黑体", 8, 4),
        "Heading 4": (11, C_MUTED, "黑体", 6, 3),
    }
    for name, (size, color, ea, before, after) in specs.items():
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        set_ea_font(h, ea)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)

    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)


def add_runs(p, text, base_size=None, base_color=None):
    """处理 **bold** 与 `code` 内联标记。"""
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        bold = code = False
        if seg.startswith("**") and seg.endswith("**"):
            seg, bold = seg[2:-2], True
        elif seg.startswith("`") and seg.endswith("`"):
            seg, code = seg[1:-1], True
        r = p.add_run(seg)
        r.bold = bold
        if code:
            r.font.name = "Consolas"
            set_ea_font(r, "Menlo")
            r.font.size = Pt((base_size or 10.5) - 0.5)
            r.font.color.rgb = RGBColor(0x0B, 0x54, 0x8A)
        elif base_size:
            r.font.size = Pt(base_size)
        if base_color and not code:
            r.font.color.rgb = base_color


def shade_para(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def shade_cell(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            text = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if ri == 0:
                shade_cell(cell, "1B2A4A")
                add_runs(p, text, base_size=9.5)
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                if ri % 2 == 0:
                    shade_cell(cell, "F2F5FA")
                add_runs(p, text, base_size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def md_to_docx(md_path: Path, out_path: Path):
    text = md_path.read_text(encoding="utf-8")
    doc = Document()
    style_setup(doc)
    # 页脚
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(f"《AI驱动的数据中台项目实战》课程资料 · 由教学方案源文件导出 · {out_path.stem}")
    fr.font.size = Pt(8)
    fr.font.color.rgb = C_MUTED

    lines = text.splitlines()
    i, n = 0, len(lines)
    in_code = False
    code_buf = []
    while i < n:
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            if in_code:
                # 输出代码块
                for cl in code_buf:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    shade_para(p, C_CODE_BG)
                    r = p.add_run(cl if cl else " ")
                    r.font.name = "Consolas"
                    set_ea_font(r, "Menlo")
                    r.font.size = Pt(8.5)
                code_buf = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if s in ("---", "***", "___"):
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            shade_para(p, "EEF3FB")
            add_runs(p, s.lstrip("> ").strip(), base_color=C_MUTED)
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            title = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2))
            h = doc.add_heading("", level=level)
            add_runs(h, title)
            i += 1
            continue
        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)[.、)]\s*(.*)$", s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(f"{m.group(1)}. ")
            r.bold = True
            r.font.color.rgb = C_ACCENT
            add_runs(p, m.group(2))
            i += 1
            continue
        p = doc.add_paragraph()
        add_runs(p, s)
        i += 1
    doc.save(out_path)


def collect_markdowns(with_project_docs=False):
    """返回 [(md_path, 导出名stem)]，导出名冲突时加父目录前缀。"""
    items = [(ROOT / "README.md", "README-课程总览")]
    for p in sorted(DOCS.rglob("*.md")):
        rel_parent = p.parent.name
        if not with_project_docs and rel_parent == "dap-v2工程文档":
            continue
        stem = p.stem
        items.append((p, stem))
    # 冲突检查：stem 重复加父目录前缀
    seen = {}
    out = []
    for path, stem in items:
        if stem in seen:
            stem = f"{path.parent.name}-{stem}"
        seen[stem] = True
        out.append((path, stem))
    return out


def convert_pdfs(docx_dir: Path, pdf_dir: Path):
    pdf_dir.mkdir(parents=True, exist_ok=True)
    files = sorted(docx_dir.glob("*.docx"))
    # LibreOffice headless 批量转换易因实例冲突/大批量挂起：
    # 每批 4 个文件、独立 profile、独立进程，逐批执行。
    import tempfile
    batch = 4
    done = 0
    for k in range(0, len(files), batch):
        chunk = files[k:k + batch]
        with tempfile.TemporaryDirectory() as td:
            res = subprocess.run(
                [SOFFICE, f"-env:UserInstallation=file://{td}",
                 "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir)]
                + [str(f) for f in chunk],
                capture_output=True, text=True, timeout=300,
            )
        ok = sum(1 for f in chunk if (pdf_dir / (f.stem + ".pdf")).exists())
        done += ok
        print(f"  [pdf] 批 {k//batch + 1}/{(len(files)+batch-1)//batch}：{ok}/{len(chunk)}")
    print(f"PDF 转换完成：{done}/{len(files)} 份")


def main():
    with_project = "--with-project-docs" in sys.argv
    skip_pdf = "--skip-pdf" in sys.argv
    WORD_OUT.mkdir(parents=True, exist_ok=True)
    PDF_OUT.mkdir(parents=True, exist_ok=True)

    items = collect_markdowns(with_project)
    print(f"待转换 {len(items)} 份 md")
    for src, stem in items:
        dst = WORD_OUT / f"{stem}.docx"
        try:
            md_to_docx(src, dst)
            print(f"  [docx] {src.relative_to(ROOT)} -> {dst.relative_to(ROOT)}")
        except Exception as e:
            print(f"  [FAIL] {src}: {e}")
    if not skip_pdf:
        convert_pdfs(WORD_OUT, PDF_OUT)
    print("\n导出完成：")
    print(f"  Word -> {WORD_OUT}")
    print(f"  PDF  -> {PDF_OUT}")


if __name__ == "__main__":
    main()
